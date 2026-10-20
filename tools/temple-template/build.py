"""Builds data/Navagraha_Temple_Data_Template.xlsx.

Usage: python3 tools/temple-template/build.py data/Navagraha_Temple_Data_Template.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

OUT = sys.argv[1]
ROWS = 200  # rows pre-formatted for entry on open-ended sheets

F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="5B2C06")
INPUT_FILL = PatternFill("solid", fgColor="FFF4CC")
PREFILL_FILL = PatternFill("solid", fgColor="EDEDED")
EX_FONT = Font(name=F, size=10, italic=True, color="808080")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

STATUS = ["To collect", "Collected", "Verified", "Needs re-check"]
SOURCE_TYPES = ["Temple office (in person)", "Temple office (phone)", "Temple official website",
                "HR&CE / Govt notice", "Temple priest", "Partner driver / guide",
                "Google Maps (dated)", "Other (explain in notes)"]
YESNO = ["Yes", "No", "Unknown"]

TEMPLES = [
    ("TEM-SUR", "Surya (Sun)", "Suriyanar Kovil", "Suriyanar Kovil", "Tamil Nadu"),
    ("TEM-CHA", "Chandra (Moon)", "Thingalur", "Thingalur", "Tamil Nadu"),
    ("TEM-ANG", "Angaraka / Sevvai (Mars)", "Vaitheeswaran Koil", "Vaitheeswaran Koil", "Tamil Nadu"),
    ("TEM-BUD", "Budha (Mercury)", "Thiruvenkadu", "Thiruvenkadu", "Tamil Nadu"),
    ("TEM-GUR", "Guru (Jupiter)", "Alangudi", "Alangudi", "Tamil Nadu"),
    ("TEM-SHU", "Shukra (Venus)", "Kanjanur", "Kanjanur", "Tamil Nadu"),
    ("TEM-SHA", "Shani (Saturn)", "Thirunallar", "Thirunallar", "Puducherry (Karaikal)"),
    ("TEM-RAH", "Rahu", "Thirunageswaram", "Thirunageswaram", "Tamil Nadu"),
    ("TEM-KET", "Ketu", "Keezhperumpallam", "Keezhperumpallam", "Tamil Nadu"),
]
HUB = ("HUB-KUM", "Kumbakonam (town centre)")
ORIGINS = [("ORG-MAA", "Chennai (city)"), ("ORG-TRZ", "Tiruchirappalli airport"),
           ("ORG-TJ", "Thanjavur junction (railway)"), ("ORG-KMU", "Kumbakonam railway station"),
           ("ORG-MV", "Mayiladuthurai junction (railway)")]

VERIFY_COLS = [("Source (who / which page)", 30), ("Source type", 24), ("Verified by", 16),
               ("Verified on", 13), ("Status", 15), ("Notes", 36)]

wb = Workbook()

def style_header(ws, headers):
    for c, (h, w) in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(name=F, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2" if ws.title != "Read me" else None

def body(ws, first, last, ncols, prefilled_cols=()):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            if cell.font is None or not cell.font.italic:
                cell.font = Font(name=F, size=10)
            cell.fill = PREFILL_FILL if c in prefilled_cols else INPUT_FILL
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)

def add_list(ws, values, col, first, last):
    dv = DataValidation(type="list", formula1=values, allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col}{first}:{col}{last}")

def lists_ref(name):
    return f"={name}"

def example_row(ws, r, values):
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = EX_FONT

def fmt(ws, col, first, last, number_format):
    for r in range(first, last + 1):
        ws[f"{col}{r}"].number_format = number_format

# ---------- Lists (hidden) ----------
lst = wb.active
lst.title = "Lists"
named = {}
def put_list(col, title, values):
    lst.cell(row=1, column=col, value=title).font = Font(name=F, bold=True)
    for i, v in enumerate(values, 2):
        lst.cell(row=i, column=col, value=v).font = Font(name=F)
    L = get_column_letter(col)
    return f"Lists!${L}$2:${L}${len(values)+1}"

temple_ids = [t[0] for t in TEMPLES]
node_ids = temple_ids + [HUB[0]] + [o[0] for o in ORIGINS]
R_STATUS = put_list(1, "Status", STATUS)
R_SRC = put_list(2, "Source type", SOURCE_TYPES)
R_YN = put_list(3, "Yes/No", YESNO)
R_TEMPLE = put_list(4, "Temple ID", temple_ids)
R_DAY = put_list(5, "Day type", ["Every day", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
                                 "Saturday", "Sunday", "Weekdays (Mon-Fri)", "Weekends (Sat-Sun)",
                                 "Tamil month first day", "Amavasai", "Pournami", "Pradosham"])
R_SESSION = put_list(6, "Session", ["Morning", "Afternoon", "Evening", "Full day"])
R_SPECIAL = put_list(7, "Special day type", ["Festival", "Graha peyarchi (transit)", "Eclipse",
                                             "Special closure", "Renovation / kumbabhishekam", "Other"])
R_EFFECT = put_list(8, "Effect", ["Closed", "Changed timings", "Extended timings", "Very heavy crowd",
                                  "Heavy crowd", "Vehicle restrictions", "Other"])
R_SVC = put_list(9, "Service type", ["Archana", "Abhishekam", "Homam", "Pariharam pooja",
                                     "Special darshan / ticket", "Other"])
R_BOOKVIA = put_list(10, "Booked via", ["Temple counter on the day", "Temple office in advance",
                                        "Temple priest", "Partner / agent", "Online (official)"])
R_NODE = put_list(11, "Place ID", node_ids)
R_RULECAT = put_list(12, "Rule category", ["State border / permit", "Vehicle entry / tax", "Parking",
                                           "Dress code", "Photography", "Footwear", "Mobile phones",
                                           "Accessibility", "Other"])
lst.sheet_state = "hidden"

# ---------- Read me ----------
rm = wb.create_sheet("Read me", 0)
rm.column_dimensions["A"].width = 3
rm.column_dimensions["B"].width = 110
lines = [
    ("Navagraha AI — Verified Temple Data Collection", "title"),
    ("Version 0.1 · 20 October 2026 · Owner: Content Administrator · Feeds FS NAV-TEM-001..003, MR-29, MR-32, MR-33, MR-47, MR-48, MR-49", "sub"),
    ("", None),
    ("Why this matters", "h"),
    ("The app will only show temple timings, services and fees that are in this workbook and marked Verified. The AI is not allowed to guess them. "
     "Blank or unverified data means the app says \"not confirmed\" instead of giving an answer.", None),
    ("", None),
    ("How to fill it in", "h"),
    ("1. Fill the YELLOW cells. GREY cells are pre-filled — check them and correct if wrong.", None),
    ("2. Every row needs a Source, Source type, Verified by and Verified on date. Use \"Temple office (in person)\" or \"(phone)\" wherever possible.", None),
    ("3. Set Status: To collect → Collected (you have it, not yet cross-checked) → Verified (confirmed from a reliable source) → Needs re-check.", None),
    ("4. Times as 24-hour HH:MM (e.g. 06:00, 20:30). Dates as DD-MM-YYYY. Money in Indian Rupees, numbers only.", None),
    ("5. One row per opening session: a temple open 06:00–12:30 and 16:00–20:30 has two rows. Different hours on Sunday? Add Sunday rows.", None),
    ("6. Rows in grey italics starting with EXAMPLE show the format only. They are not real data — overwrite or delete them.", None),
    ("7. If two sources disagree, add a row for each and explain in Notes. Operations will decide (FS EC-50).", None),
    ("8. Religious meaning (e.g. what a pariharam is believed to do) goes in the \"Traditional belief\" column, worded as belief — never as a medical or guaranteed claim (FS MR-43).", None),
    ("", None),
    ("Sheets", "h"),
    ("Temples — one row per Navagraha temple: address, location, access, visit duration.", None),
    ("Schedules — opening hours by day and session.", None),
    ("Special days — festivals, graha peyarchi, eclipses, closures for the next 12 months.", None),
    ("Services — pooja / pariharam / archana catalogue with temple fees and booking rules.", None),
    ("Travel times — drive time between every pair of temples and from common starting points.", None),
    ("Entry rules — Thirunallar / Puducherry entry, vehicle tax, dress code, photography and similar rules.", None),
    ("Progress — automatic count of rows per status. No input needed.", None),
]
for i, (t, kind) in enumerate(lines, 2):
    c = rm.cell(row=i, column=2, value=t)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title":
        c.font = Font(name=F, size=16, bold=True, color="5B2C06")
    elif kind == "sub":
        c.font = Font(name=F, size=9, italic=True, color="606060")
    elif kind == "h":
        c.font = Font(name=F, size=12, bold=True)
    else:
        c.font = Font(name=F, size=10)
leg = len(lines) + 3
rm.cell(row=leg, column=2, value="Legend").font = Font(name=F, size=12, bold=True)
for j, (fill, txt) in enumerate([(INPUT_FILL, "Yellow — please fill in"),
                                 (PREFILL_FILL, "Grey — pre-filled, please check")], 1):
    c = rm.cell(row=leg + j, column=2, value=txt)
    c.fill = fill; c.font = Font(name=F, size=10); c.border = BORDER

# ---------- Temples ----------
ws = wb.create_sheet("Temples")
hdr = [("Temple ID", 11), ("Graha (planet)", 20), ("Commonly known as", 20), ("Official temple name", 30),
       ("Main deity", 22), ("Town / village", 18), ("District", 16), ("State", 18), ("Full address", 34),
       ("PIN code", 10), ("Latitude", 11), ("Longitude", 11), ("Google Maps link", 26), ("Temple office phone", 16),
       ("Official website", 24), ("Walk from parking to sanctum (metres)", 14), ("Steps / stairs to climb", 16),
       ("Wheelchair access possible?", 13), ("Resting / seating available?", 13), ("Clean toilets nearby?", 12),
       ("Typical visit time — normal day (min)", 14), ("Typical visit time — busy day (min)", 14),
       ("Parking notes", 26), ("Senior traveller notes", 30)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
for i, (tid, graha, known, town, state) in enumerate(TEMPLES, 2):
    for c, v in zip((1, 2, 3, 6, 8), (tid, graha, known, town, state)):
        ws.cell(row=i, column=c, value=v)
    ws.cell(row=i, column=n - 1, value="To collect")
body(ws, 2, 10, n, prefilled_cols=(1, 2, 3, 6, 8))
for col in ("R", "S", "T"):
    add_list(ws, R_YN, col, 2, 10)
add_list(ws, R_SRC, get_column_letter(n - 4), 2, 10)
add_list(ws, R_STATUS, get_column_letter(n - 1), 2, 10)
fmt(ws, get_column_letter(n - 2), 2, 10, "DD-MM-YYYY")
fmt(ws, "K", 2, 10, "0.000000"); fmt(ws, "L", 2, 10, "0.000000")
ws["P1"].comment = Comment("Approximate walking distance in metres from where the vehicle parks to the main shrine. Important for Senior-Friendly plans (NAV-SEN-004).", "Navagraha AI")
ws["U1"].comment = Comment("From entering the temple to leaving, on an ordinary weekday, without special pooja.", "Navagraha AI")

# ---------- Schedules ----------
ws = wb.create_sheet("Schedules")
hdr = [("Temple ID", 11), ("Day type", 20), ("Session", 12), ("Opens (HH:MM)", 11), ("Closes (HH:MM)", 11),
       ("Last entry (HH:MM)", 11), ("Valid from", 13), ("Valid to (blank = no end)", 13)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
example_row(ws, 2, ["EXAMPLE TEM-XXX", "Every day", "Morning", "06:00", "12:30", "12:15", "01-11-2026", "",
                    "Temple office, Mr. ___ (phone)", "Temple office (phone)", "Ravi", "18-10-2026", "Verified",
                    "EXAMPLE — format only, not real data"])
body(ws, 3, ROWS, n)
add_list(ws, R_TEMPLE, "A", 3, ROWS); add_list(ws, R_DAY, "B", 3, ROWS); add_list(ws, R_SESSION, "C", 3, ROWS)
add_list(ws, R_SRC, get_column_letter(n - 4), 3, ROWS); add_list(ws, R_STATUS, get_column_letter(n - 1), 3, ROWS)
for col in ("D", "E", "F"): fmt(ws, col, 3, ROWS, "@")
for col in ("G", "H", get_column_letter(n - 2)): fmt(ws, col, 3, ROWS, "DD-MM-YYYY")

# ---------- Special days ----------
ws = wb.create_sheet("Special days")
hdr = [("Temple ID", 11), ("Name of day / event", 26), ("Type", 20), ("From date", 13), ("To date", 13),
       ("Effect", 18), ("Changed opens (HH:MM)", 11), ("Changed closes (HH:MM)", 11),
       ("Expected extra wait (min)", 12), ("Recurs every year?", 11)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
example_row(ws, 2, ["EXAMPLE TEM-XXX", "Sani peyarchi", "Graha peyarchi (transit)", "DD-MM-YYYY", "DD-MM-YYYY",
                    "Very heavy crowd", "04:00", "22:00", "180", "No",
                    "Temple notice board", "Temple office (in person)", "Ravi", "18-10-2026", "Collected",
                    "EXAMPLE — format only, not real data"])
body(ws, 3, ROWS, n)
add_list(ws, R_TEMPLE, "A", 3, ROWS); add_list(ws, R_SPECIAL, "C", 3, ROWS); add_list(ws, R_EFFECT, "F", 3, ROWS)
add_list(ws, R_YN, "J", 3, ROWS)
add_list(ws, R_SRC, get_column_letter(n - 4), 3, ROWS); add_list(ws, R_STATUS, get_column_letter(n - 1), 3, ROWS)
for col in ("G", "H"): fmt(ws, col, 3, ROWS, "@")
for col in ("D", "E", get_column_letter(n - 2)): fmt(ws, col, 3, ROWS, "DD-MM-YYYY")

# ---------- Services ----------
ws = wb.create_sheet("Services")
hdr = [("Service ID", 11), ("Temple ID", 11), ("Service name (as the temple calls it)", 28), ("Type", 18),
       ("Traditional belief / purpose (worded as belief)", 34), ("Available days", 18),
       ("Time window (HH:MM–HH:MM)", 16), ("Duration (min)", 10), ("Max people per booking", 10),
       ("Book how many days ahead?", 11), ("Booked via", 22), ("Temple fee (₹)", 11), ("Fee is per", 14),
       ("Items devotee must bring", 24), ("Items available at temple", 24), ("Dress requirement", 20)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
example_row(ws, 2, ["EXAMPLE SVC-001", "TEM-XXX", "Archana", "Archana", "Devotees traditionally offer this to …",
                    "Every day", "06:30–12:00", "15", "6", "0", "Temple counter on the day", "50", "Per person",
                    "None", "Flowers, coconut", "Traditional dress preferred",
                    "Temple counter price board", "Temple office (in person)", "Ravi", "18-10-2026", "Collected",
                    "EXAMPLE — format only, not real data"])
body(ws, 3, ROWS, n)
add_list(ws, R_TEMPLE, "B", 3, ROWS); add_list(ws, R_SVC, "D", 3, ROWS); add_list(ws, R_BOOKVIA, "K", 3, ROWS)
add_list(ws, '"Per person,Per family / group,Per booking"', "M", 3, ROWS)
add_list(ws, R_SRC, get_column_letter(n - 4), 3, ROWS); add_list(ws, R_STATUS, get_column_letter(n - 1), 3, ROWS)
fmt(ws, "L", 3, ROWS, "#,##0"); fmt(ws, get_column_letter(n - 2), 3, ROWS, "DD-MM-YYYY")
ws["L1"].comment = Comment("Fee paid to the temple on site. Shown to the customer separately from our prices (MR-33).", "Navagraha AI")

# ---------- Travel times ----------
ws = wb.create_sheet("Travel times")
hdr = [("From ID", 10), ("From", 26), ("To ID", 10), ("To", 26), ("Distance (km)", 10),
       ("Normal drive (min)", 10), ("Peak drive (min)", 10), ("Peak hours / days", 18),
       ("Same both ways?", 10), ("Road notes (condition, narrow roads, level crossings)", 32)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
nodes = [(t[0], t[2]) for t in TEMPLES] + [HUB]
pairs = [(a, b) for i, a in enumerate(nodes) for b in nodes[i + 1:]] + [(o, HUB) for o in ORIGINS]
for r, (a, b) in enumerate(pairs, 2):
    ws.cell(row=r, column=1, value=a[0]); ws.cell(row=r, column=2, value=a[1])
    ws.cell(row=r, column=3, value=b[0]); ws.cell(row=r, column=4, value=b[1])
    ws.cell(row=r, column=9, value="Yes")
    ws.cell(row=r, column=n - 1, value="To collect")
last_pair = len(pairs) + 1
body(ws, 2, last_pair + 30, n, prefilled_cols=(1, 2, 3, 4))
for r in range(last_pair + 1, last_pair + 31):
    for c in (1, 2, 3, 4): ws.cell(row=r, column=c).fill = INPUT_FILL
add_list(ws, R_NODE, "A", last_pair + 1, last_pair + 30); add_list(ws, R_NODE, "C", last_pair + 1, last_pair + 30)
add_list(ws, R_YN, "I", 2, last_pair + 30)
add_list(ws, R_SRC, get_column_letter(n - 4), 2, last_pair + 30)
add_list(ws, R_STATUS, get_column_letter(n - 1), 2, last_pair + 30)
fmt(ws, get_column_letter(n - 2), 2, last_pair + 30, "DD-MM-YYYY")
ws["I1"].comment = Comment("Yes if the drive is about the same in both directions. If No, add a separate row for the return direction in the blank rows at the bottom.", "Navagraha AI")
ws["F1"].comment = Comment("Door-to-door drive in an ordinary car at a quiet time. Use a driver's experience or Google Maps at a stated date/time.", "Navagraha AI")

# ---------- Entry rules ----------
ws = wb.create_sheet("Entry rules")
hdr = [("Temple ID (blank = applies to whole trip)", 14), ("Category", 20), ("Rule", 44),
       ("Applies to", 20), ("Cost (₹, if any)", 11), ("What the driver / customer must do", 36)] + VERIFY_COLS
style_header(ws, hdr)
n = len(hdr)
example_row(ws, 2, ["EXAMPLE TEM-XXX", "Dress code", "Shorts not allowed inside", "All visitors", "",
                    "Tell travellers at booking and in T-48h checklist",
                    "Temple board at entrance", "Temple office (in person)", "Ravi", "18-10-2026", "Collected",
                    "EXAMPLE — format only, not real data"])
prompts = [("TEM-SHA", "State border / permit", "Puducherry (Karaikal) entry — any permit or check needed for tourist vehicles?"),
           ("TEM-SHA", "Vehicle entry / tax", "Puducherry entry tax / road tax for Tamil Nadu-registered tourist vehicles — amount and where paid"),
           ("", "Vehicle entry / tax", "Any other entry fees, tolls or parking charges on the circuit")]
for r, (tid, cat, rule) in enumerate(prompts, 3):
    ws.cell(row=r, column=1, value=tid); ws.cell(row=r, column=2, value=cat); ws.cell(row=r, column=3, value=rule)
    ws.cell(row=r, column=n - 1, value="To collect")
body(ws, 3, 60, n)
add_list(ws, R_TEMPLE, "A", 3, 60); add_list(ws, R_RULECAT, "B", 3, 60)
add_list(ws, R_SRC, get_column_letter(n - 4), 3, 60); add_list(ws, R_STATUS, get_column_letter(n - 1), 3, 60)
fmt(ws, "E", 3, 60, "#,##0"); fmt(ws, get_column_letter(n - 2), 3, 60, "DD-MM-YYYY")

# ---------- Progress ----------
pg = wb.create_sheet("Progress", 1)
pg.column_dimensions["A"].width = 18
for c in "BCDEFG": pg.column_dimensions[c].width = 15
pg["A1"] = "Collection progress (automatic)"; pg["A1"].font = Font(name=F, size=14, bold=True, color="5B2C06")
pg["A2"] = "Counts rows by Status. Grey EXAMPLE rows are not counted."; pg["A2"].font = Font(name=F, size=9, italic=True, color="606060")
heads = ["Sheet"] + STATUS + ["Total rows", "% Verified"]
for c, h in enumerate(heads, 1):
    cell = pg.cell(row=4, column=c, value=h)
    cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR_FILL; cell.border = BORDER
sheets = [("Temples", "AC", 10), ("Schedules", "M", ROWS), ("Special days", "O", ROWS),
          ("Services", "U", ROWS), ("Travel times", "O", last_pair + 30), ("Entry rules", "K", 60)]
for i, (s, col, last) in enumerate(sheets, 5):
    first = 2 if s in ("Temples", "Travel times") else 3
    rng = f"'{s}'!${col}${first}:${col}${last}"
    pg.cell(row=i, column=1, value=s)
    for j, st in enumerate(STATUS, 2):
        pg.cell(row=i, column=j, value=f'=COUNTIF({rng},"{st}")')
    pg.cell(row=i, column=6, value=f"=SUM(B{i}:E{i})")
    pg.cell(row=i, column=7, value=f"=IF(F{i}=0,0,D{i}/F{i})")
    pg.cell(row=i, column=7).number_format = "0%"
    for c in range(1, 8):
        pg.cell(row=i, column=c).font = Font(name=F, size=10); pg.cell(row=i, column=c).border = BORDER
tr = 5 + len(sheets)
pg.cell(row=tr, column=1, value="All sheets")
for c in range(2, 7):
    L = get_column_letter(c)
    pg.cell(row=tr, column=c, value=f"=SUM({L}5:{L}{tr-1})")
pg.cell(row=tr, column=7, value=f"=IF(F{tr}=0,0,D{tr}/F{tr})").number_format = "0%"
for c in range(1, 8):
    pg.cell(row=tr, column=c).font = Font(name=F, size=10, bold=True); pg.cell(row=tr, column=c).border = BORDER

wb.save(OUT)
print("saved", OUT, "pairs", len(pairs))
